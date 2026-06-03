"""
data/export.py — Excel / CSV 匯出邏輯，與 GUI 解耦。
"""
import csv
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from config import ROW_FIELDS


def save_excel(rounds: list, path: str, log_cb):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hdr_fill  = PatternFill("solid", fgColor="1e2a38")
    hdr_font  = Font(name="Consolas", bold=True, color="00c8ff")
    hdr_align = Alignment(horizontal="center")

    for idx, rdata in enumerate(rounds):
        if not rdata:
            continue
        ws = wb.create_sheet(title=f"Round_{idx + 1:02d}")
        for ci, h in enumerate(ROW_FIELDS, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = hdr_align
        for ri, row in enumerate(rdata, 2):
            for ci, h in enumerate(ROW_FIELDS, 1):
                ws.cell(row=ri, column=ci, value=row.get(h, ""))
        for col in ws.columns:
            w = max(len(str(c.value)) if c.value else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(w + 2, 24)

    wb.save(path)
    total = sum(len(r) for r in rounds)
    log_cb(f"[SAVE] Excel → {path}  ({len(rounds)} 輪, {total} 筆)")


def save_csv(data: list, path: str, log_cb):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=ROW_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(data)
    log_cb(f"[SAVE] CSV → {path}  ({len(data)} 筆)")


def make_default_filename(state, n_rounds: int, ext: str) -> str:
    K, B, M, *_ = state.get_params()
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"impedance_K{K:.3f}_B{B:.4f}_M{M:.4f}_{n_rounds}rounds_{ts}.{ext}"
